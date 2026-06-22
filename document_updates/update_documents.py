from pathlib import Path
from shutil import copy2
from copy import deepcopy
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path("/Users/rachit/Code/Voice Vision/document_updates")
IEEE_SOURCE = Path("/Users/rachit/Library/Containers/net.whatsapp.WhatsApp/Data/tmp/documents/D6D5E699-B36A-44A7-BA53-95AE671BCF9B/IEEE Paper Final.docx")
REPORT_SOURCE = Path("/Users/rachit/Library/Containers/net.whatsapp.WhatsApp/Data/tmp/documents/8FF40923-AB26-4101-8B58-62F67E2CA512/EL REPORT.docx")
IEEE_OUTPUT = ROOT / "Second Sight - IEEE Paper Final.docx"
REPORT_OUTPUT = ROOT / "Second Sight - EL Report.docx"

SKILL_SCRIPTS = Path("/Users/rachit/.codex/plugins/cache/openai-primary-runtime/documents/26.619.11828/skills/documents/scripts")
sys.path.insert(0, str(SKILL_SCRIPTS))
from table_geometry import apply_table_geometry, column_widths_from_weights, section_content_width_dxa


def find_paragraph(doc, prefix):
    for paragraph in doc.paragraphs:
        if paragraph.text.strip().startswith(prefix):
            return paragraph
    raise ValueError(f"Paragraph not found: {prefix}")


def set_paragraph_text(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(text)


def insert_after(anchor, text, style=None):
    paragraph = anchor._parent.add_paragraph()
    if style is not None:
        paragraph.style = style
    paragraph.add_run(text)
    anchor._p.addnext(paragraph._p)
    return paragraph


def insert_after_with_format(anchor, text, format_source):
    paragraph = anchor._parent.add_paragraph()
    if paragraph._p.pPr is not None:
        paragraph._p.remove(paragraph._p.pPr)
    if format_source._p.pPr is not None:
        paragraph._p.insert(0, deepcopy(format_source._p.pPr))
    paragraph.add_run(text)
    anchor._p.addnext(paragraph._p)
    return paragraph


def add_bullet_paragraph(doc, text, num_id="7"):
    paragraph = doc.add_paragraph(text, style="List Paragraph")
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), num_id)
    num_pr.append(ilvl)
    num_pr.append(num_id_element)
    p_pr.append(num_pr)
    return paragraph


def move_table_after(table, anchor):
    anchor._p.addnext(table._tbl)
    return table._tbl


def replace_project_name_everywhere(doc):
    replacements = {
        "VoiceVision": "Second Sight",
        "Voice Vision": "Second Sight",
    }
    for text_node in doc.element.body.iter(qn("w:t")):
        if text_node.text:
            for old, new in replacements.items():
                text_node.text = text_node.text.replace(old, new)

    for section in doc.sections:
        for part in (section.header, section.footer):
            for text_node in part._element.iter(qn("w:t")):
                if text_node.text:
                    for old, new in replacements.items():
                        text_node.text = text_node.text.replace(old, new)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def format_comparison_table(table, section):
    table.style = "Table Grid"
    total_width = section_content_width_dxa(section)
    widths = column_widths_from_weights([1.25, 1.0, 1.0], total_width)
    apply_table_geometry(table, widths)

    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = 0
                paragraph.paragraph_format.space_after = 0
                if row_index == 0:
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.bold = True
        if row_index == 0:
            for cell in row.cells:
                shade_cell(cell, "D9EAF7")


def update_ieee():
    copy2(IEEE_SOURCE, IEEE_OUTPUT)
    doc = Document(IEEE_OUTPUT)

    set_paragraph_text(find_paragraph(doc, "VoiceVision- Assistance for Visually Impaired"),
        "Second Sight - Assistance for Visually Impaired")

    set_paragraph_text(find_paragraph(doc, "Blind and visually impaired (BVI) individuals face"),
        "Blind and visually impaired (BVI) individuals face daily challenges in safe navigation, understanding nearby obstacles, reading printed text, recognizing currency, and performing secure financial transactions without external assistance. This paper presents Second Sight, an iOS-based assistive application that combines ARKit camera capture, Apple Vision and Core ML obstacle detection, filtered OCR, TensorFlow Lite currency recognition, and secure QR-based payments. The updated currency module uses a MobileNetV2-based six-class Indian banknote classifier with OCR support, explicit note-presence gating, confidence-margin rejection, and temporal consensus. The text-reading pipeline rejects keyboard fragments and lexical noise before offering user-controlled speech. The AI-powered Describe feature sends the current camera image through a backend proxy to MiniCPM-V 4.6 in LM Studio and speaks a concise scene summary. QR Pay Mode requires allowlisted QR payloads, manual amount entry, and Face ID authorization. A Node.js/Express backend with MongoDB persistence, optional Telegram notifications, and a React dashboard complete the modular assistive platform.")

    set_paragraph_text(find_paragraph(doc, "Keywords:"),
        "Keywords: Assistive technology, visual impairment, ARKit, Core ML, TensorFlow Lite, MobileNetV2, Vision OCR, currency recognition, MiniCPM-V, QR scanning, Face ID, digital wallet.")

    set_paragraph_text(find_paragraph(doc, "To address these challenges"),
        "To address these challenges, this paper proposes Second Sight, a unified iOS assistive system for real-time navigation support, contextual text reading, Indian currency identification, optional vision-language scene description, and secure wallet transfers. The system performs on-device YOLO obstacle detection, filters OCR output before user-controlled reading, and recognizes six Indian banknote denominations using a TensorFlow Lite classifier combined with OCR and rejection logic. QR payments are restricted to allowlisted codes and require Face ID authorization. The remainder of this paper describes the architecture, methodology, implementation, measured model comparison, limitations, and future enhancements of Second Sight.")

    set_paragraph_text(find_paragraph(doc, "VoiceVision is developed as an integrated"),
        "Second Sight is an integrated platform consisting of an iOS SwiftUI application, a backend service, and a web dashboard. ARKit publishes camera frames to mode-specific perception modules. YOLO obstacle detection runs through Vision and Core ML, text extraction uses Vision OCR with semantic filtering, and Indian currency classification runs through TensorFlow Lite. QR scanning supports secure payment initiation. The named Describe feature captures the current frame, compresses it as an image, and sends it to a backend perception endpoint. The backend forwards the multimodal request to MiniCPM-V 4.6 hosted in LM Studio on a networked Mac, and the returned concise scene summary is spoken through the app. Mode-based execution limits competing workloads and excessive speech.")

    set_paragraph_text(find_paragraph(doc, "Text recognition is implemented using TextRecognitionService.swift"),
        "Text recognition is implemented in TextRecognitionService.swift using Vision OCR in accurate mode with language correction. Observations below 0.55 confidence, very small regions, fragmented single-character sequences, keyboard-row patterns, and lexically unrecognizable text are rejected. Useful signs, numbers, prices, and common abbreviations remain eligible. The system announces that meaningful text is available and waits for a tap; the selected text is validated again immediately before speech.")

    architecture_ocr = find_paragraph(doc, "Text recognition is implemented in TextRecognitionService.swift")
    architecture_currency_heading = find_paragraph(doc, "Currency Recognition Module")
    describe_heading = insert_after_with_format(
        architecture_ocr,
        "AI-Powered Describe Feature",
        architecture_currency_heading,
    )
    insert_after(describe_heading,
        "Describe provides semantic scene understanding beyond fixed object labels. When invoked, MiniCPMService captures the current ARKit frame, converts it to a compressed JPEG, and sends it with a navigation-focused prompt to the backend /api/perception/analyze endpoint. The backend uses the OpenAI-compatible LM Studio API to run the vision-capable MiniCPM-V 4.6 model. The response is reduced to a concise description of important objects, layout, and navigation-relevant context, then delivered through AVSpeechSynthesizer. Unlike obstacle and currency inference, Describe is not fully on-device and requires the Mac, backend, LM Studio server, and loaded model to be reachable on the same network.")

    set_paragraph_text(find_paragraph(doc, "Currency recognition is implemented through CurrencyRecognitionService.swift"),
        "Currency recognition is implemented in CurrencyRecognitionService.swift using a MobileNetV2-based TensorFlow Lite classifier and OCR parsing. The model supports ₹10, ₹20, ₹50, ₹100, ₹200, and ₹500 notes; ₹2000 is not included because that class is absent from the training dataset. A Vision rectangle request first checks for a banknote-like object. Classification is accepted only above 0.85 confidence, with at least a 0.25 lead over the second-ranked class, and after two consecutive matching frames. Brightness gating and OCR evidence provide additional safeguards. Currency mode announces the denomination only and does not mutate wallet balance.")

    set_paragraph_text(find_paragraph(doc, "The VoiceVision iOS application is implemented"),
        "The Second Sight iOS application is implemented using SwiftUI. ContentView.swift coordinates camera input, recognition modes, gestures, and speech. ARKit frames are selectively routed to Core ML obstacle detection, filtered Vision OCR, TensorFlow Lite currency recognition, QR scanning, and the optional backend-proxied MiniCPM-V description service. This modular mode-based design maintains responsive processing and understandable feedback.")

    set_paragraph_text(find_paragraph(doc, "Text recognition is implemented using Apple Vision OCR"),
        "Text recognition uses VNRecognizeTextRequest in accurate mode with language correction. A 0.55 confidence threshold, minimum region area, and observation-level semantic checks suppress OCR noise. The filter rejects keyboard rows such as QWERTY/ASDF, heavily fragmented single-letter sequences, and dictionary-unrecognized jargon while preserving meaningful signs, monetary values, and common abbreviations. Second Sight prompts before reading and repeats the filter when the user taps, preventing stale or noisy text from reaching AVSpeechSynthesizer.")

    set_paragraph_text(find_paragraph(doc, "When the user taps the screen after a text prompt"),
        "When the user taps after a text prompt, Second Sight validates the assembled text again and uses SpeechService.swift with AVSpeechSynthesizer to read only accepted content. Cooldowns prevent repeated prompts, and OCR pauses during speech to avoid overlapping announcements. This two-stage validation is particularly important for keyboards and control panels, where OCR may otherwise concatenate individual keys into meaningless speech.")

    methodology_ocr = find_paragraph(doc, "When the user taps after a text prompt")
    methodology_currency_heading = next(
        p for p in doc.paragraphs if p.text.strip() == "Currency Recognition"
    )
    methodology_describe_heading = insert_after_with_format(
        methodology_ocr,
        "D.  AI-Powered Describe Feature Implementation",
        methodology_currency_heading,
    )
    methodology_describe_heading._p.get_or_add_pPr().remove(
        methodology_describe_heading._p.get_or_add_pPr().numPr
    )
    insert_after(methodology_describe_heading,
        "The Describe workflow is implemented through MiniCPMService. The service throttles scene requests, encodes the latest CVPixelBuffer as a compressed image, and submits a mode-specific prompt and optional OCR context to the backend. The backend is configured for the LM Studio OpenAI-compatible chat-completions endpoint using the minicpm-v-4.6 model identifier. On success, the app stores the returned summary and speaks it, while duplicate summaries and rapid repeat requests are suppressed by cooldown logic.")

    for heading, label in (
        (methodology_currency_heading, "E.  Currency Recognition"),
        (find_paragraph(doc, "QR Pay Mode and Secure Transfer Implementation"), "F.  QR Pay Mode and Secure Transfer Implementation"),
        (find_paragraph(doc, "Backend, Database, Telegram, and Dashboard Implementation"), "G.  Backend, Database, Telegram, and Dashboard Implementation"),
    ):
        set_paragraph_text(heading, label)
        heading._p.get_or_add_pPr().remove(heading._p.get_or_add_pPr().numPr)

    set_paragraph_text(find_paragraph(doc, "Currency recognition in VoiceVision is implemented"),
        "Currency recognition in Second Sight uses the trained IndianCurrency.tflite model through a pinned TensorFlow Lite Swift runtime. The model receives a 224 × 224 float32 RGB image and returns probabilities for six denominations: ₹10, ₹20, ₹50, ₹100, ₹200, and ₹500. OCR independently extracts rupee symbols and denomination digits. The previous 299 × 299 Core ML classifier, which included ₹2000, was removed after comparative testing showed substantially lower accuracy on the shared held-out benchmark.")

    set_paragraph_text(find_paragraph(doc, "To ensure stable recognition during continuous scanning"),
        "To prevent a closed-set classifier from assigning a denomination to an empty table or unrelated object, the pipeline first requires a banknote-like rectangle or explicit OCR currency evidence. Model-only decisions require confidence ≥0.85 and a top-two probability margin ≥0.25. Two consecutive frames must agree; changing predictions reset consensus, and frames without note evidence clear pending detections. Brightness gating remains active. A double tap enables currency mode, and confirmed denominations are spoken without changing wallet balance.")

    set_paragraph_text(find_paragraph(doc, "VoiceVision was evaluated through controlled experiments"),
        "Second Sight was evaluated through controlled functional tests and an offline currency-model comparison. Both currency classifiers were tested on the same held-out set of 293 images spanning ₹10, ₹20, ₹50, ₹100, ₹200, and ₹500 notes. The dataset was checked for exact train-test duplicates and none were found. Additional device tests covered empty scenes, keyboard OCR, mode switching, QR validation, Face ID gating, backend updates, dashboard synchronization, and the Describe request path. LM Studio was verified to expose MiniCPM-V 4.6 and accept image input through its OpenAI-compatible endpoint.")

    set_paragraph_text(find_paragraph(doc, "The OCR subsystem demonstrated strong capability"),
        "The revised OCR subsystem reduced irrelevant speech by filtering at both recognition and tap-to-read stages. Accurate Vision OCR, language correction, a 0.55 confidence threshold, minimum-area filtering, keyboard-sequence rejection, fragmented-character rejection, and lexical checks prevent key rows and random visual patterns from being read aloud. Useful signs, prices, numbers, and common abbreviations remain available through the user-controlled prompt-and-tap interaction.")

    heading = find_paragraph(doc, "Currency Recognition Stability and Practical Usability")
    set_paragraph_text(heading, "Currency Model Comparison, Stability, and Practical Usability")
    result_para = find_paragraph(doc, "Currency recognition benefited from the hybrid OCR")
    set_paragraph_text(result_para,
        "On the common 293-image test set, the previous Core ML model correctly classified 47 images (16.0%), whereas the trained TensorFlow Lite model correctly classified 267 images (91.1%). New-model class accuracy was 86.8% for ₹10, 93.3% for ₹20, 97.7% for ₹50, 79.1% for ₹100, 96.1% for ₹200, and 92.9% for ₹500. The older model’s corresponding scores were 1.9%, 3.3%, 4.5%, 2.3%, 0.0%, and 97.6%. Because this benchmark comes from the new model’s six-class dataset and excludes ₹2000, the result demonstrates suitability for the current deployment rather than universal superiority across every capture distribution.")
    insert_after(result_para,
        "The deployment was further hardened with note-shape detection, probability-margin rejection, two-frame consecutive consensus, active clearing of empty scenes, and OCR corroboration. These controls address the classifier’s closed-set behavior: without an explicit background class, a softmax classifier always ranks some denomination even when no note is present.", result_para.style)

    set_paragraph_text(find_paragraph(doc, "During evaluation, it was observed that backend base URL"),
        "During evaluation, backend addressing was found to differ between simulator and physical-device operation: an iPhone requires the Mac’s LAN address rather than localhost. The Describe feature requires the backend, LM Studio, and the loaded MiniCPM-V 4.6 vision model to remain reachable on the same network; if those services are unavailable, on-device obstacle, OCR, and currency features continue to operate. The current currency model covers six denominations and has no explicit background or ₹2000 class; geometric, confidence, OCR, and temporal rejection reduce this limitation, while future retraining should add representative non-currency images. Distance estimation remains approximate without depth sensing, and older devices may experience thermal or frame-rate constraints.")

    set_paragraph_text(find_paragraph(doc, "This paper presented VoiceVision"),
        "This paper presented Second Sight, an assistive iOS system integrating navigation feedback, meaningful text reading, Indian currency recognition, the AI-powered Describe feature, and secure QR wallet transfers. Describe uses MiniCPM-V 4.6 through LM Studio to convert a live camera frame into a concise spoken scene summary. ARKit supplies camera frames; Vision and Core ML perform obstacle detection; Vision OCR uses semantic filtering; and the trained TensorFlow Lite currency classifier achieved 91.1% accuracy on the shared 293-image benchmark compared with 16.0% for the removed Core ML model. Note-presence gating and consecutive consensus reduce empty-scene false positives, while the filtered tap-to-read workflow suppresses keyboard and gibberish output. Face ID-gated QR payments, MongoDB persistence, Telegram alerts, and dashboard monitoring complete the system.")

    replace_project_name_everywhere(doc)
    doc.save(IEEE_OUTPUT)


def update_report():
    copy2(REPORT_SOURCE, REPORT_OUTPUT)
    doc = Document(REPORT_OUTPUT)

    set_paragraph_text(find_paragraph(doc, "This project presents a multi-component"),
        "This project presents Second Sight, a multi-component assistive system consisting of an iOS application, a Node.js backend with persistent storage, and a React dashboard. The system assists visually impaired users through obstacle awareness, meaningful text reading, Indian currency recognition, the AI-powered Describe feature, and secure QR-based transactions with Face ID authorization. Describe sends the current camera image through the backend to MiniCPM-V 4.6 in LM Studio and speaks the returned scene summary. On-device perception uses ARKit, Vision, Core ML, and TensorFlow Lite.")

    set_paragraph_text(find_paragraph(doc, "The iOS application continuously captures camera frames"),
        "The iOS application continuously captures ARKit camera frames and routes them according to the active mode. YOLO obstacle detection runs through Core ML and Vision. Text recognition uses accurate Vision OCR with confidence, geometry, keyboard-pattern, fragmentation, and lexical filtering. Currency recognition combines OCR with a trained six-class TensorFlow Lite classifier. The Describe feature compresses the current camera frame, sends it through /api/perception/analyze, and speaks the concise summary generated by MiniCPM-V 4.6 through the LM Studio backend connection.")

    set_paragraph_text(find_paragraph(doc, "The iOS application was developed using SwiftUI"),
        "The iOS application was developed using SwiftUI and ARKit. Vision and Core ML handle obstacle detection, Vision performs filtered OCR and QR scanning, and TensorFlow Lite runs the Indian currency classifier. A centralized view coordinates mode switching, gesture handling, speech prompts, MiniCPM scene-description requests, QR scanning, amount entry, and Face ID authentication.")

    report_ios_paragraph = find_paragraph(doc, "The iOS application was developed using SwiftUI")
    report_backend_heading = find_paragraph(doc, "5.2.2 Backend Development")
    report_describe_heading = insert_after_with_format(
        report_ios_paragraph,
        "5.2.2 AI-Powered Describe Feature",
        report_backend_heading,
    )
    insert_after(report_describe_heading,
        "Describe extends Second Sight beyond fixed-class detection by explaining the complete visible scene. MiniCPMService captures the current ARKit frame, converts it to a compressed JPEG, and sends it with a concise accessibility-focused prompt to the backend perception endpoint. The backend forwards the multimodal request to the OpenAI-compatible LM Studio server running MiniCPM-V 4.6. The returned summary is spoken using AVSpeechSynthesizer. Request cooldowns and duplicate-summary suppression prevent repetitive output. This feature requires the Mac backend and LM Studio to be reachable over the local network; the other on-device recognition features remain available independently.")
    set_paragraph_text(report_backend_heading, "5.2.3 Backend Development")
    set_paragraph_text(find_paragraph(doc, "5.2.3 Frontend Dashboard Development"), "5.2.4 Frontend Dashboard Development")
    set_paragraph_text(find_paragraph(doc, "5.2.4 Testing"), "5.2.5 Testing")

    set_paragraph_text(find_paragraph(doc, "Testing was conducted across all components"),
        "Testing covered model accuracy, empty-scene currency rejection, meaningful-text filtering, keyboard OCR rejection, mode switching, QR validation, biometric authorization, backend endpoints, and dashboard behavior. The old and new currency models were evaluated on the same 293-image held-out set, and exact train-test duplicate checking found no overlap.")

    set_paragraph_text(find_paragraph(doc, "The iOS application was able to capture live camera frames"),
        "The iOS application reliably captured and processed live camera frames. Obstacle detection provided navigation-oriented audio feedback. Text recognition now uses accurate OCR and semantic filtering to reject keyboard rows, fragmented characters, and unrecognizable jargon before prompting; a second validation occurs when the user taps to read. This prevents irrelevant sequences such as Q-W-E-R-T-Y from being spoken while preserving signs, prices, numbers, and common abbreviations. The Describe feature successfully routed image prompts to the configured MiniCPM-V 4.6 endpoint and returned scene summaries for spoken delivery when the Mac services were available.")

    set_paragraph_text(find_paragraph(doc, "Currency recognition was tested under normal lighting"),
        "Currency recognition was evaluated on 293 held-out images across ₹10, ₹20, ₹50, ₹100, ₹200, and ₹500. The removed Core ML model achieved 16.0% accuracy (47/293), while the trained TensorFlow Lite model achieved 91.1% (267/293). The deployed pipeline additionally requires note-shape evidence, ≥0.85 confidence, a ≥0.25 top-two margin, and two consecutive matching frames. The model does not support ₹2000 because that class was absent from its dataset. QR Pay Mode continued to enforce allowlisting, mode exclusivity, and Face ID authorization.")

    set_paragraph_text(find_paragraph(doc, "The iOS application serves as the core prototype"),
        "The iOS application is the core Second Sight prototype and uses the device camera for real-time perception. It integrates Core ML obstacle detection, semantically filtered Vision OCR, TensorFlow Lite currency recognition, QR scanning, and the named Describe feature. Describe uses backend-proxied MiniCPM-V 4.6 to turn the current image into a spoken semantic scene summary. Controlled modes and cooldowns limit workload conflicts and unnecessary audio output.")

    set_paragraph_text(find_paragraph(doc, "Initially, the focus was on developing the iOS application’s core"),
        "Development began with ARKit camera integration, Core ML obstacle detection, and Vision OCR. Currency recognition was subsequently migrated from the original Core ML classifier to a trained MobileNetV2-based TensorFlow Lite model. The pipeline was then hardened with note-presence detection, confidence-margin rejection, and true consecutive-frame consensus. OCR was upgraded with language correction and semantic filtering, and MiniCPM-V 4.6 was configured through LM Studio for optional scene descriptions.")

    set_paragraph_text(find_paragraph(doc, "The iOS mobile application was tested for correct camera frame"),
        "The iOS app was tested for camera acquisition, mode switching, device builds, and real-time processing. Currency tests included the shared 293-image benchmark plus bare-table scenarios to verify that empty scenes clear detections. Text tests included keyboards and fragmented labels to verify that meaningless sequences are rejected before speech. QR tests covered allowlisted and untrusted payloads, amount confirmation, and successful, cancelled, and failed Face ID flows. Backend and dashboard validation covered balance updates, transaction recording, error handling, and polling.")

    set_paragraph_text(find_paragraph(doc, "The iOS application provides real-time assistance"),
        "The Second Sight iOS application provides obstacle awareness, filtered text reading, six-denomination currency recognition, the AI-powered Describe feature, and secure QR wallet payments. Describe sends a live camera image through the backend to MiniCPM-V 4.6 in LM Studio and speaks the returned scene summary. Core ML remains responsible for on-device obstacle detection, while currency inference uses TensorFlow Lite and OCR uses Apple Vision. The currency model comparison showed 91.1% accuracy for the deployed model versus 16.0% for the removed model on the common held-out set. Face ID remains mandatory for QR payments.")

    # Correct the Android-oriented tool table so it matches the implemented iOS stack.
    tools_table = doc.tables[3]
    tools = [
        ("Swift and SwiftUI", "iOS application and accessible interface"),
        ("ARKit", "Continuous camera capture and spatial session"),
        ("Vision and Core ML", "OCR, QR scanning, and YOLO11n obstacle detection"),
        ("TensorFlow Lite", "MobileNetV2 Indian currency classification"),
        ("AVSpeechSynthesizer", "Spoken prompts and recognized-content output"),
        ("LM Studio / MiniCPM-V 4.6", "Optional local vision-language scene description"),
        ("Node.js, MongoDB, React", "Wallet backend, persistence, and dashboard"),
        ("Xcode and devicectl", "iPhone builds, signing, installation, and testing"),
    ]
    while len(tools_table.rows) < len(tools) + 1:
        tools_table.add_row()
    while len(tools_table.rows) > len(tools) + 1:
        tools_table._tbl.remove(tools_table.rows[-1]._tr)
    tools_table.cell(0, 0).text = "Tool"
    tools_table.cell(0, 1).text = "Purpose"
    for index, (tool, purpose) in enumerate(tools, start=1):
        tools_table.cell(index, 0).text = tool
        tools_table.cell(index, 1).text = purpose

    anchor = find_paragraph(doc, "All primary objectives of the project were successfully achieved")
    heading = insert_after(anchor, "7.3 Currency Model Comparison and Implemented Improvements", "Heading 2")
    intro = insert_after(heading,
        "Both currency models were evaluated using identical preprocessing on the same held-out set of 293 images. The set contained six denominations and no exact duplicates of training images. The comparison therefore measures suitability for the current six-class deployment; it does not evaluate ₹2000 or every possible real-world background.")

    table = doc.add_table(rows=1, cols=3)
    table.rows[0].cells[0].text = "Criterion"
    table.rows[0].cells[1].text = "Old model"
    table.rows[0].cells[2].text = "New model"
    comparison_rows = [
        ("Runtime and input", "Core ML, 299 × 299 BGR", "TensorFlow Lite, 224 × 224 RGB float32"),
        ("Classes", "7, including ₹2000", "6: ₹10 to ₹500; no ₹2000"),
        ("Correct predictions", "47 of 293", "267 of 293"),
        ("Overall accuracy", "16.0%", "91.1%"),
    ]
    for criterion, old, new in comparison_rows:
        cells = table.add_row().cells
        cells[0].text = criterion
        cells[1].text = old
        cells[2].text = new
    format_comparison_table(table, doc.sections[0])
    move_table_after(table, intro)
    cursor = table._tbl

    per_class = doc.add_paragraph(
        "New-model accuracy by denomination was ₹10: 86.8%, ₹20: 93.3%, ₹50: 97.7%, ₹100: 79.1%, ₹200: 96.1%, and ₹500: 92.9%. The old model scored 1.9%, 3.3%, 4.5%, 2.3%, 0.0%, and 97.6%, respectively."
    )
    cursor.addnext(per_class._p)
    cursor = per_class._p

    improvements = [
        "Replaced IndianCurrency.mlmodel with the trained IndianCurrency.tflite model and integrated a pinned TensorFlow Lite Swift runtime.",
        "Added banknote-shape detection, brightness gating, 0.85 confidence, a 0.25 top-two probability margin, and two consecutive matching frames before announcing currency.",
        "Corrected consensus logic so the first prediction is no longer published immediately; empty scenes clear pending results.",
        "Upgraded Vision OCR to accurate mode with language correction and 0.55 confidence filtering.",
        "Added keyboard-row, fragmented-character, dictionary, signage, number, and acronym checks at recognition time and again before speech.",
        "Configured optional MiniCPM-V 4.6 scene description through LM Studio and the backend; this feature requires the Mac service to be reachable, unlike on-device currency inference.",
    ]
    for item in improvements:
        paragraph = add_bullet_paragraph(doc, item)
        cursor.addnext(paragraph._p)
        cursor = paragraph._p

    replace_project_name_everywhere(doc)

    # The source template carries many empty trailing paragraphs after the QR image.
    # Once added content moves the QR page, those spacers create a blank final page.
    qr_image_paragraph = next(
        paragraph for index, paragraph in enumerate(doc.paragraphs)
        if index > 145 and paragraph._p.xpath('.//w:drawing')
    )
    remove_after = False
    for paragraph in list(doc.paragraphs):
        if remove_after:
            paragraph._element.getparent().remove(paragraph._element)
        elif paragraph._p is qr_image_paragraph._p:
            remove_after = True

    doc.save(REPORT_OUTPUT)


if __name__ == "__main__":
    ROOT.mkdir(parents=True, exist_ok=True)
    update_ieee()
    update_report()
    print(IEEE_OUTPUT)
    print(REPORT_OUTPUT)
